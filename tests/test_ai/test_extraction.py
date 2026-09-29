from io import BytesIO

from docx import Document as DocxDocument
from pypdf import PdfWriter

from modules.ai.extraction import extract_text


def test_extract_text_plain_decodes_utf8():
    assert extract_text(file_bytes=b"hello world", mime_type="text/plain") == (
        "hello world"
    )


def test_extract_text_markdown_uses_plain_decode_path():
    content = b"# Heading\n\nSome body text."
    assert extract_text(file_bytes=content, mime_type="text/markdown") == (
        content.decode()
    )


def test_extract_text_docx_extracts_paragraphs():
    doc = DocxDocument()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    buffer = BytesIO()
    doc.save(buffer)

    text = extract_text(
        file_bytes=buffer.getvalue(),
        mime_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )

    assert text == "First paragraph.\nSecond paragraph."


def test_extract_text_pdf_with_no_text_layer_returns_empty_string():
    """A scanned/image-only PDF has no extractable text — pypdf can't OCR,
    so this must come back empty rather than raise, since that's exactly
    the case that produces zero chunks downstream (not a bug)."""
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = BytesIO()
    writer.write(buffer)

    text = extract_text(file_bytes=buffer.getvalue(), mime_type="application/pdf")

    assert text == ""
